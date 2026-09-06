import docx, sys
from pathlib import Path

files = [
    r'task1\task1报告new(1).docx',
    r'task2\报告\task2测试结果分析报告.docx',
    r'task3\task3实验报告.docx',
]
out = []
for fp in files:
    out.append(f'\n========== {fp} ==========')
    d = docx.Document(fp)
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            out.append(t)
    for tbl_i, tbl in enumerate(d.tables, 1):
        out.append(f'[Table {tbl_i}]')
        for row in tbl.rows:
            out.append(' | '.join(c.text.strip() for c in row.cells))

with open('_experiments.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('lines:', len(out))
