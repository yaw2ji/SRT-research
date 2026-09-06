from pptx import Presentation
from pypdf import PdfReader

out = []
prs = Presentation('内容护栏-0812.pptx')
for i, slide in enumerate(prs.slides, 1):
    out.append(f'\n--- Slide {i} ---')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = ''.join(run.text for run in para.runs).strip()
                if text:
                    out.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                out.append(' | '.join(cells))

out.append('\n===== PDF =====')
reader = PdfReader('AntiLLMRe_paper.pdf')
for i, page in enumerate(reader.pages, 1):
    out.append(f'\n--- Page {i} ---')
    out.append(page.extract_text() or '')

with open('_extracted.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('done', len(out))
